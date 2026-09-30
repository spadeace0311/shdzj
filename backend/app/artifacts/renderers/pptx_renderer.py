from __future__ import annotations

import hashlib
import os
import re
from collections.abc import Mapping
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches

from app.artifacts.catalog import load_catalog
from app.artifacts.domain import ArtifactNameContext, ProductionMode
from app.artifacts.naming import build_artifact_file_name
from app.artifacts.renderers.base import RenderResult
from app.artifacts.renderers.docx_renderer import DocumentImage, DocumentRenderSpec
from app.artifacts.template_builder import DECISION_TEMPLATE_SLIDES
from app.config import settings

PPTX_RENDERER_VERSION = "artifact-pptx-renderer-v1"
PLACEHOLDER_PATTERN = re.compile(r"\{\{[^{}]*\}\}")
DECISION_TEMPLATE_FILE = "decision-template.pptx"


class PptxRenderer:
    def __init__(self, *, template_path: str | Path | None = None) -> None:
        self._template_path = Path(
            template_path
            or Path(settings.artifact_template_root) / DECISION_TEMPLATE_FILE
        )

    async def render(
        self,
        spec: DocumentRenderSpec,
        output_path: Path,
    ) -> RenderResult:
        if spec.document_type != "pptx":
            raise ValueError("PptxRenderer only accepts pptx document specs")

        target = Path(output_path)
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_name(f".{target.name}.{os.getpid()}.pptx")
        actual_template_checksum = _sha256_path(spec.template_path)
        if actual_template_checksum != spec.template_checksum:
            raise ValueError("template checksum mismatch")

        file_name = build_artifact_file_name(
            ArtifactNameContext(
                place=str(spec.event.get("place") or "未命名震中"),
                magnitude=float(spec.event.get("magnitude", 0.0)),
                display_name=spec.display_name,
                version=1,
                generated_at=datetime.now(UTC),
                production_mode=ProductionMode(spec.production_mode),
            ),
            extension="pptx",
        )

        presentation = Presentation(str(spec.template_path))
        self._validate_template(presentation, spec)

        values = dict(spec.control_fields)
        values["file_name"] = file_name
        values["version"] = spec.template_version
        values["generated_by"] = "地震应急辅助决策系统"
        self._replace_text_placeholders(presentation, values)

        images = {image.key: image for image in spec.images}
        self._replace_image_placeholders(presentation, images)

        unresolved_count = self._unresolved_placeholder_count(presentation)
        if unresolved_count:
            unresolved_texts = self._unresolved_placeholder_texts(presentation)
            raise ValueError(
                f"rendered PPTX contains {unresolved_count} unresolved placeholders: "
                f"{unresolved_texts}"
            )

        presentation.save(temporary)
        os.replace(temporary, target)

        checksum = _sha256_path(target)
        image_checksums = self._image_checksums(images)
        render_manifest: dict[str, Any] = {
            "marker": spec.marker or "",
            "template_version": spec.template_version,
            "template_checksum": spec.template_checksum,
            "image_checksums": image_checksums,
            "image_paths": {
                key: str(image.path)
                for key, image in images.items()
            },
            "images": list(images),
            "page_count": len(presentation.slides),
            "unresolved_placeholder_count": unresolved_count,
            "renderer_version": PPTX_RENDERER_VERSION,
            "context_fingerprint": spec.context_fingerprint,
            "input_fingerprint": spec.input_fingerprint,
            "quality": spec.quality.to_dict(),
            "control_fields": values,
        }
        return RenderResult(
            path=target,
            format="pptx",
            width=1280,
            height=720,
            dpi=96,
            checksum=checksum,
            quality=spec.quality,
            task_status=(
                "degraded"
                if spec.quality.needs_review
                or spec.quality.degradation_reasons
                else "succeeded"
            ),
            file_name=file_name,
            render_manifest=render_manifest,
            non_empty_ratio=1.0,
            size_bytes=target.stat().st_size,
            generated_at=datetime.now(UTC),
            page_count=len(presentation.slides),
            control_fields=dict(spec.control_fields),
        )

    def _validate_template(
        self,
        presentation: Any,
        spec: DocumentRenderSpec,
    ) -> None:
        if len(presentation.slides) != len(DECISION_TEMPLATE_SLIDES):
            raise ValueError(
                "decision template must contain exactly "
                f"{len(DECISION_TEMPLATE_SLIDES)} slides"
            )
        for slide, (layout_key, layout_label, _body, image_keys) in zip(
            presentation.slides,
            DECISION_TEMPLATE_SLIDES,
        ):
            if not any(
                shape.has_text_frame
                and shape.name == f"decision-{layout_key}"
                and shape.text_frame.text == layout_label
                for shape in slide.shapes
            ):
                raise ValueError(
                    f"decision template layout {layout_key!r} is missing"
                )
            actual_image_keys = {
                shape.name
                for shape in slide.shapes
                if shape.name.startswith("decision-image-")
            }
            expected_image_keys = {
                f"decision-image-{key}" for key in image_keys
            }
            if actual_image_keys != expected_image_keys:
                raise ValueError(
                    f"decision template image slots for {layout_key!r} are invalid"
                )
        definition = load_catalog(
            settings.artifact_catalog_path
        ).get(spec.artifact_key, "a3v-professional")
        template_text = _presentation_text(presentation)
        missing = [
            field
            for field in definition.control_fields
            if not field.startswith("image:")
            and f"{{{{{field}}}}}" not in template_text
        ]
        if missing:
            raise ValueError(
                "decision template is missing declared text placeholders: "
                + ", ".join(missing)
            )

    def _replace_text_placeholders(
        self,
        presentation: Any,
        values: Mapping[str, Any],
    ) -> None:
        for slide in presentation.slides:
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    if "{{" not in paragraph.text:
                        continue
                    text = paragraph.text
                    for key, value in values.items():
                        text = text.replace(
                            f"{{{{{key}}}}}",
                            "" if value is None else str(value),
                        )
                    if text == paragraph.text:
                        continue
                    if paragraph.runs:
                        paragraph.runs[0].text = text
                        for run in paragraph.runs[1:]:
                            run.text = ""
                    else:
                        paragraph.add_run().text = text

    def _replace_image_placeholders(
        self,
        presentation: Any,
        images: Mapping[str, DocumentImage],
    ) -> None:
        for slide in presentation.slides:
            shapes = list(slide.shapes)
            for image_key, image in images.items():
                placeholder = next(
                    (
                        shape
                        for shape in shapes
                        if shape.has_text_frame
                        and f"{{{{image:{image_key}}}}}" in shape.text_frame.text
                    ),
                    None,
                )
                if placeholder is None:
                    continue
                self._validate_image(image)
                self._insert_image(slide, placeholder, image)
                _remove_shape(placeholder)

    def _validate_image(self, image: DocumentImage) -> None:
        path = image.path
        if not path.is_absolute():
            raise ValueError(f"PPTX image must use an absolute local path: {path}")
        if not path.is_file():
            raise FileNotFoundError(f"PPTX image does not exist: {path}")
        actual_checksum = _sha256_path(path)
        if image.checksum and image.checksum != actual_checksum:
            raise ValueError(f"PPTX image checksum mismatch for {image.key}")

    def _insert_image(
        self,
        slide: Any,
        placeholder: Any,
        image: DocumentImage,
    ) -> None:
        left = placeholder.left if placeholder.left is not None else Inches(0.45)
        top = placeholder.top if placeholder.top is not None else Inches(4.05)
        width = (
            placeholder.width
            if placeholder.width is not None
            else Inches(3.82)
        )
        height = (
            placeholder.height
            if placeholder.height is not None
            else Inches(2.72)
        )
        slide.shapes.add_picture(
            str(image.path),
            left,
            top,
            width,
            height,
        )

    def _unresolved_placeholder_count(self, presentation: Any) -> int:
        return sum(
            len(PLACEHOLDER_PATTERN.findall(shape.text_frame.text))
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.has_text_frame
            and PLACEHOLDER_PATTERN.search(shape.text_frame.text)
        )

    def _unresolved_placeholder_texts(self, presentation: Any) -> list[str]:
        return [
            shape.text_frame.text
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.has_text_frame
            and PLACEHOLDER_PATTERN.search(shape.text_frame.text)
        ]

    def _image_checksums(
        self,
        images: Mapping[str, DocumentImage],
    ) -> dict[str, str]:
        checksums: dict[str, str] = {}
        for key, image in images.items():
            checksum = image.checksum or _sha256_path(image.path)
            if image.checksum and checksum != image.checksum:
                raise ValueError(f"image checksum mismatch for {key}")
            checksums[key] = checksum
        return checksums


def _remove_shape(shape: Any) -> None:
    parent = shape._element.getparent()
    if parent is not None:
        parent.remove(shape._element)


def _presentation_text(presentation: Any) -> str:
    return "\n".join(
        shape.text_frame.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )


def _sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while True:
            chunk = source.read(1024 * 1024)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()
