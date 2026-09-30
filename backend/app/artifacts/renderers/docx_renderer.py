from __future__ import annotations

import hashlib
import os
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from PIL import Image

from app.artifacts.catalog import load_catalog
from app.artifacts.context import DocumentRenderContext, FrozenAssetVersion
from app.artifacts.domain import ArtifactNameContext, ProductionMode
from app.artifacts.naming import build_artifact_file_name
from app.artifacts.renderers.base import RenderQuality, RenderResult
from app.artifacts.renderers.chart_renderer import ChartRenderer, ChartSpec
from app.config import settings

DOCX_RENDERER_VERSION = "artifact-docx-renderer-v1"
PLACEHOLDER_PATTERN = re.compile(r"\{\{[^{}]*\}\}")
BACKGROUND_DOC_KEYS = {
    "doc.background",
    "doc.housing",
    "doc.economy",
    "doc.population",
    "doc.key_targets",
    "doc.spatial_distances",
    "doc.area_overview",
    "doc.historical_catalog",
}


@dataclass(frozen=True, slots=True)
class DocumentImage:
    key: str
    path: Path
    checksum: str | None = None
    width: int | None = None
    height: int | None = None
    caption: str = ""


@dataclass(frozen=True, slots=True)
class DocumentSection:
    title: str
    body: tuple[str, ...] = ()
    table: tuple[tuple[str, ...], ...] = ()
    image_keys: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class DocumentRenderSpec:
    artifact_key: str
    display_name: str
    title: str
    document_type: str
    template_path: Path
    template_version: str
    template_checksum: str
    context_fingerprint: str
    input_fingerprint: str
    production_mode: str
    marker: str | None
    control_fields: Mapping[str, Any]
    sections: tuple[DocumentSection, ...]
    images: tuple[DocumentImage, ...]
    charts: tuple[ChartSpec, ...] = ()
    quality: RenderQuality = field(default_factory=RenderQuality)
    event: Mapping[str, Any] = field(default_factory=dict)
    asset_versions: Mapping[str, FrozenAssetVersion] = field(default_factory=dict)

    def __post_init__(self) -> None:
        if not self.artifact_key.strip():
            raise ValueError("artifact_key must not be empty")
        if not self.display_name.strip():
            raise ValueError("display_name must not be empty")
        if not self.title.strip():
            raise ValueError("document title must not be empty")
        if not self.template_path.is_file():
            raise FileNotFoundError(f"document template does not exist: {self.template_path}")
        if len(self.template_checksum) != 64:
            raise ValueError("template checksum must be a SHA-256 hex digest")
        if len(self.context_fingerprint) != 64:
            raise ValueError("context_fingerprint must be a SHA-256 hex digest")
        if len(self.input_fingerprint) != 64:
            raise ValueError("input_fingerprint must be a SHA-256 hex digest")
        if not self.sections:
            raise ValueError("document spec must contain at least one section")
        image_keys = [image.key for image in self.images]
        if len(image_keys) != len(set(image_keys)):
            raise ValueError("document image keys must be unique")


def _sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while True:
            chunk = source.read(1024 * 1024)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _isoformat(value: Any) -> str:
    if value is None:
        return "不适用"
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def _t1_display(value: Any) -> str:
    if value is None:
        return "不适用（人工/测试/演练）"
    return f"T1: {_isoformat(value)}"


def _mode_marker(production_mode: str) -> str | None:
    return {
        "live": None,
        "manual": None,
        "test": "【测试】",
        "drill": "【演练】",
        "replay": "【测试回放】",
    }.get(production_mode)


def _product_payload(manifest: Mapping[str, Any], product_key: str) -> Mapping[str, Any]:
    assessment = manifest.get("assessment")
    if isinstance(assessment, Mapping):
        products = assessment.get("products")
        if isinstance(products, Mapping) and isinstance(products.get(product_key), Mapping):
            return dict(products[product_key])
    loss = manifest.get("loss")
    if isinstance(loss, Mapping):
        products = loss.get("products")
        if isinstance(products, Mapping) and isinstance(products.get(product_key), Mapping):
            return dict(products[product_key])
    return {}


def _metric(payload: Mapping[str, Any], *keys: str) -> str:
    for key in keys:
        value = payload.get(key)
        if value is not None and str(value).strip():
            return str(value)
    return "数据缺失，待复核"


def _asset_summary(
    asset_versions: Mapping[str, FrozenAssetVersion],
    asset_key: str,
) -> str:
    asset = asset_versions.get(asset_key)
    if asset is None or asset.checksum is None:
        return f"{asset_key}: 数据缺失，待复核"
    return f"{asset_key}: 已冻结 / 版本 {asset.version or '未知'} / {asset.checksum[:12]}"


def build_background_document_spec(
    context: DocumentRenderContext,
    artifact_key: str,
) -> DocumentRenderSpec:
    if artifact_key not in BACKGROUND_DOC_KEYS:
        raise ValueError(f"unsupported background document key: {artifact_key}")
    if context.artifact_key != artifact_key:
        raise ValueError("document context artifact_key does not match requested artifact")

    catalog = load_catalog(settings.artifact_catalog_path)
    definition = catalog.get(artifact_key, context.output_profile)
    manifest = dict(context.manifest)
    event = dict(manifest.get("event") or {})
    t1_at = event.get("t1_at", manifest.get("t1_at"))
    production_mode = context.production_mode
    marker = context.marker or _mode_marker(production_mode)

    template = context.template_versions.get("background-template")
    template_version = (
        str(template.get("version"))
        if isinstance(template, Mapping) and template.get("version")
        else "v1"
    )
    template_path = Path(settings.artifact_template_root) / "background-template.docx"
    template_checksum = (
        str(template.get("checksum"))
        if isinstance(template, Mapping) and template.get("checksum")
        else _sha256_path(template_path)
    )
    if len(template_checksum) != 64:
        template_checksum = _sha256_path(template_path)

    asset_versions = dict(context.asset_versions)
    missing_required = [
        key
        for key in definition.required_assets
        if asset_versions.get(key) is None
        or asset_versions[key].checksum is None
    ]
    missing_optional = [
        key
        for key in definition.optional_assets
        if asset_versions.get(key) is None
        or asset_versions[key].checksum is None
    ]
    degradation_reasons = [
        f"{key} 数据缺失" for key in missing_required + missing_optional
    ]
    needs_review = bool(missing_required or missing_optional)

    input_payload: dict[str, Any] = {
        "artifact_key": artifact_key,
        "output_profile": context.output_profile,
        "context_fingerprint": context.context_fingerprint,
        "template_version": template_version,
        "template_checksum": template_checksum,
        "assets": {
            key: {
                "checksum": asset.checksum,
                "version": asset.version,
            }
            for key, asset in sorted(asset_versions.items())
        },
    }
    input_fingerprint = _sha256_bytes(
        str(input_payload).encode("utf-8")
    )

    control_fields = {
        "event_name": event.get("place") or "未命名震中",
        "magnitude": f"{float(event.get('magnitude', 0.0)):.1f} 级",
        "origin_time": _isoformat(event.get("origin_time")),
        "longitude": f"{float(event.get('longitude', 0.0)):.6f}",
        "latitude": f"{float(event.get('latitude', 0.0)):.6f}",
        "depth_km": f"{float(event.get('depth_km', 0.0)):.2f} km",
        "t1_at": _t1_display(t1_at),
        "data_source": "离线冻结数据资产 / 评估产品 / 系统受控模板",
        "report_type": definition.display_name,
        "mode_marker": marker or "正式",
        "generated_at": datetime.now(UTC).isoformat(),
        "artifact_version": "V001",
        "revision_no": str(event.get("revision_no") or manifest.get("event", {}).get("revision_no") or "1"),
        "assessment_run_no": (
            str(manifest.get("assessment", {}).get("assessment_run_id"))
            if isinstance(manifest.get("assessment"), Mapping)
            else "不适用"
        ),
        "data_asset_snapshot_fingerprint": (
            str(manifest.get("assessment", {}).get("data_asset_snapshot_fingerprint"))
            if isinstance(manifest.get("assessment"), Mapping)
            else "未冻结"
        ),
        "model_parameter_version": (
            str(manifest.get("loss", {}).get("parameter_package", {}).get("version"))
            if isinstance(manifest.get("loss"), Mapping)
            and isinstance(manifest["loss"].get("parameter_package"), Mapping)
            else "未冻结"
        ),
        "template_version": template_version,
        "quality_grade": definition.quality_policy,
        "degradation_reasons": "；".join(degradation_reasons) or "无",
        "needs_review": "待复核" if needs_review else "无",
        "version": template_version,
        "file_name": "",
        "page_number": "",
        "generated_by": "地震应急辅助决策系统",
        "title": f"{marker or ''}{definition.display_name}",
    }

    sections = _background_sections(artifact_key, manifest, asset_versions)
    if sections:
        first = sections[0]
        control_body = (
            *first.body,
            f"事件名称：{control_fields['event_name']}",
            f"数据来源：{control_fields['data_source']}",
            f"T1: {control_fields['t1_at'].removeprefix('T1: ')}",
        )
        sections = (
            DocumentSection(
                title=first.title,
                body=control_body,
                table=first.table,
                image_keys=first.image_keys,
            ),
            *sections[1:],
        )
    images = _background_images(context, artifact_key)

    return DocumentRenderSpec(
        artifact_key=artifact_key,
        display_name=definition.display_name,
        title=f"{marker or ''}{definition.display_name}",
        document_type=context.document_type,
        template_path=template_path,
        template_version=template_version,
        template_checksum=template_checksum,
        context_fingerprint=context.context_fingerprint,
        input_fingerprint=input_fingerprint,
        production_mode=production_mode,
        marker=marker,
        control_fields=control_fields,
        sections=sections,
        images=images,
        charts=_background_charts(artifact_key, manifest),
        quality=RenderQuality(
            grade=definition.quality_policy,
            needs_review=needs_review,
            degradation_reasons=tuple(degradation_reasons),
        ),
        event=event,
        asset_versions=asset_versions,
    )


def _background_images(
    context: DocumentRenderContext,
    artifact_key: str,
) -> tuple[DocumentImage, ...]:
    if artifact_key != "doc.area_overview":
        return ()
    path = context.artifact_paths.get("map.epicenter")
    if path is None or not path.is_file():
        return ()
    checksum = context.asset_versions.get("map.epicenter")
    width, height = _image_dimensions(path)
    return (
        DocumentImage(
            key="map.epicenter",
            path=path,
            checksum=checksum.checksum if checksum is not None else _sha256_path(path),
            width=width,
            height=height,
            caption="震中位置分布图",
        ),
    )


def _background_charts(
    artifact_key: str,
    manifest: Mapping[str, Any],
) -> tuple[ChartSpec, ...]:
    if artifact_key == "doc.population":
        population = _product_payload(manifest, "loss.population")
        labels = ("常住人口", "流动人口", "受灾人口")
        values = (
            _optional_float(population.get("resident")),
            _optional_float(population.get("floating")),
            _optional_float(population.get("affected")),
        )
        return (
            ChartSpec(
                chart_id="population-impact",
                title="人口影响概览",
                labels=labels,
                series=(_chart_series("人口影响", values),),
                y_label="数量",
            ),
        )
    if artifact_key == "doc.housing":
        housing = _product_payload(manifest, "loss.buildings")
        labels = ("房屋总量", "轻度破坏", "中度破坏", "严重破坏")
        values = (
            _optional_float(housing.get("total")),
            _optional_float(housing.get("slight")),
            _optional_float(housing.get("moderate")),
            _optional_float(housing.get("severe")),
        )
        return (
            ChartSpec(
                chart_id="housing-damage",
                title="房屋破坏统计",
                labels=labels,
                series=(_chart_series("房屋数量", values),),
                y_label="数量",
            ),
        )
    if artifact_key == "doc.economy":
        economy = _product_payload(manifest, "loss.economic")
        labels = ("第一产业", "第二产业", "第三产业", "经济损失")
        values = (
            _optional_float(economy.get("primary")),
            _optional_float(economy.get("secondary")),
            _optional_float(economy.get("tertiary")),
            _optional_float(economy.get("loss")),
        )
        return (
            ChartSpec(
                chart_id="economy-loss",
                title="经济结构与损失",
                labels=labels,
                series=(_chart_series("金额", values),),
                y_label="金额",
            ),
        )
    return ()


def _chart_series(name: str, values: tuple[float, ...]) -> Any:
    from app.artifacts.renderers.chart_renderer import ChartSeries

    return ChartSeries(name=name, values=values)


def _optional_float(value: Any) -> float:
    if value is None:
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _background_sections(
    artifact_key: str,
    manifest: Mapping[str, Any],
    asset_versions: Mapping[str, FrozenAssetVersion],
) -> tuple[DocumentSection, ...]:
    event = dict(manifest.get("event") or {})
    buildings = _product_payload(manifest, "loss.buildings")
    population = _product_payload(manifest, "loss.population")
    economy = _product_payload(manifest, "loss.economic")
    intensity = _product_payload(manifest, "intensity.fusion")
    historical = _product_payload(manifest, "historical.earthquakes")

    if artifact_key == "doc.background":
        return (
            DocumentSection(
                "事件参数",
                (
                    f"事件名称：{event.get('place') or '未命名震中'}",
                    f"震级：{float(event.get('magnitude', 0.0)):.1f} 级",
                    f"发震时刻：{_isoformat(event.get('origin_time'))}",
                    (
                        "震中坐标："
                        f"{float(event.get('longitude', 0.0)):.6f} E, "
                        f"{float(event.get('latitude', 0.0)):.6f} N"
                    ),
                    f"震源深度：{float(event.get('depth_km', 0.0)):.2f} km",
                ),
            ),
            DocumentSection("所在行政区", ("上海市及邻近行政区", "行政区边界数据为离线冻结资产")),
            DocumentSection(
                "历史地震",
                (
                    _asset_summary(asset_versions, "shanghai.historical.earthquakes"),
                    _metric(historical, "summary") if historical else "目录数据由历史地震资产提供",
                ),
            ),
            DocumentSection(
                "邻近断裂",
                (_asset_summary(asset_versions, "shanghai.fault"),),
            ),
            DocumentSection(
                "基础地理背景",
                ("上海市地处长江三角洲，地势低平，河网密布。", "基础地理信息来自离线行政区与底图资产。"),
            ),
        )
    if artifact_key == "doc.housing":
        return (
            DocumentSection("街镇房屋总量", (_metric(buildings, "total", "town_totals"),)),
            DocumentSection(
                "结构类型",
                ("结构分类数据来自 shanghai.building.town 冻结资产。",),
            ),
            DocumentSection(
                "破坏统计",
                (
                    f"轻度破坏：{_metric(buildings, 'slight')}",
                    f"中度破坏：{_metric(buildings, 'moderate')}",
                    f"严重破坏：{_metric(buildings, 'severe')}",
                ),
            ),
            DocumentSection(
                "覆盖与质量",
                (
                    _asset_summary(asset_versions, "shanghai.building.town"),
                    "评估产品版本与数据资产校验和以冻结快照为准。",
                ),
            ),
        )
    if artifact_key == "doc.economy":
        return (
            DocumentSection("GDP", (_metric(economy, "gdp", "total"),)),
            DocumentSection(
                "产业结构",
                (
                    f"第一产业：{_metric(economy, 'primary')}",
                    f"第二产业：{_metric(economy, 'secondary')}",
                    f"第三产业：{_metric(economy, 'tertiary')}",
                ),
            ),
            DocumentSection("经济损失", (_metric(economy, "loss", "economic_loss"),)),
            DocumentSection(
                "来源与情景",
                ("经济损失来源：loss.economic 评估产品。", "情景口径以冻结损失参数与区域配置为准。"),
            ),
        )
    if artifact_key == "doc.population":
        return (
            DocumentSection("常住人口", (_metric(population, "resident"),)),
            DocumentSection("流动人口", (_metric(population, "floating"),)),
            DocumentSection("家庭户", (_metric(population, "households"),)),
            DocumentSection("年龄结构", (_metric(population, "age_structure"),)),
            DocumentSection("受灾人口", (_metric(population, "affected"),)),
        )
    if artifact_key == "doc.key_targets":
        rows = (
            ("避难场所", _asset_summary(asset_versions, "shanghai.shelter.emergency")),
            ("学校", _asset_summary(asset_versions, "shanghai.education.school")),
            ("医院", _asset_summary(asset_versions, "shanghai.health.hospital")),
            ("危险源", _asset_summary(asset_versions, "shanghai.hazard_source")),
            ("救援队伍", _asset_summary(asset_versions, "shanghai.rescue_team")),
            ("文物单位", _asset_summary(asset_versions, "shanghai.cultural_relic")),
            ("重点目标", _asset_summary(asset_versions, "shanghai.key_target")),
        )
        return tuple(
            DocumentSection(label, (summary,))
            for label, summary in rows
        ) + (DocumentSection("重点目标列表", ("重点目标清单以冻结资产为准。",)),)
    if artifact_key == "doc.spatial_distances":
        return (
            DocumentSection("城市距离", ("上海主城区距离为空间分析结果。",)),
            DocumentSection("区县距离", ("各区县中心距离由冻结参考点计算。",)),
            DocumentSection("街镇距离", ("街镇中心距离由行政区划资产计算。",)),
            DocumentSection("主要城市距离", (_asset_summary(asset_versions, "shanghai.distance.reference_points"),)),
            DocumentSection("重点目标距离", ("重点目标距离引用重点目标资产。",)),
            DocumentSection("断裂距离", (_asset_summary(asset_versions, "shanghai.fault"),)),
        )
    if artifact_key == "doc.area_overview":
        return (
            DocumentSection(
                "地理概况",
                ("上海市地势低平，水网密集，中心城区人口与设施高度集中。",),
            ),
            DocumentSection("行政区", ("上海市及周边行政区，专业成果仅覆盖上海市域。",)),
            DocumentSection("烈度", (_metric(intensity, "summary", "grade"),)),
            DocumentSection("人口", (_metric(population, "affected", "resident"),)),
            DocumentSection("房屋", (_metric(buildings, "total", "severe"),)),
            DocumentSection("经济", (_metric(economy, "gdp", "loss"),)),
            DocumentSection(
                "关键风险",
                ("人口密集区、重点目标、危险源和断裂带为关键风险控制项。",),
                image_keys=("map.epicenter",),
            ),
        )
    if artifact_key == "doc.historical_catalog":
        return (
            DocumentSection("检索半径", ("半径：50 km（受控目录参数）",)),
            DocumentSection("震级阈值", ("震级阈值：M >= 3.0",)),
            DocumentSection(
                "历史地震目录",
                (_asset_summary(asset_versions, "shanghai.historical.earthquakes"),),
            ),
            DocumentSection("灾害地震目录", ("灾害地震目录由历史地震资产筛选得到。",)),
            DocumentSection("统计", ("记录数、最大震级和最近事件以冻结目录为准。",)),
        )
    raise ValueError(f"unsupported background document key: {artifact_key}")


class DocxRenderer:
    def __init__(
        self,
        *,
        template_path: str | Path | None = None,
        chart_renderer: ChartRenderer | None = None,
    ) -> None:
        self._template_path = Path(
            template_path or Path(settings.artifact_template_root) / "background-template.docx"
        )
        self._chart_renderer = chart_renderer or ChartRenderer()

    async def render(self, spec: DocumentRenderSpec, output_path: Path) -> RenderResult:
        target = Path(output_path)
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_name(f".{target.name}.{os.getpid()}.docx")
        document = Document(str(spec.template_path))

        self._replace_placeholders(document, spec)
        self._remove_section_placeholders(document)
        self._append_sections(document, spec)
        self._embed_images(document, spec, target.parent)
        self._set_footer(document, spec)
        self._assert_no_placeholders(document)

        page_count = self._page_count(document)
        document.save(temporary)
        os.replace(temporary, target)

        checksum = _sha256_path(target)
        image_checksums = self._resolved_image_checksums(spec, target.parent)
        render_manifest: dict[str, Any] = {
            "marker": spec.marker,
            "template_version": spec.template_version,
            "template_checksum": spec.template_checksum,
            "image_checksums": image_checksums,
            "image_paths": {
                key: str(image.path)
                for key, image in self._all_images(spec, target.parent).items()
            },
            "control_fields": dict(spec.control_fields),
            "quality": spec.quality.to_dict(),
            "page_count": page_count,
            "unresolved_placeholder_count": 0,
            "renderer_version": DOCX_RENDERER_VERSION,
            "context_fingerprint": spec.context_fingerprint,
            "input_fingerprint": spec.input_fingerprint,
        }
        file_name = build_artifact_file_name(
            ArtifactNameContext(
                place=str(spec.event.get("place") or "未命名震中"),
                magnitude=float(spec.event.get("magnitude", 0.0)),
                display_name=spec.display_name,
                version=1,
                generated_at=datetime.now(UTC),
                production_mode=ProductionMode(spec.production_mode),
            ),
            extension="docx",
        )
        return RenderResult(
            path=target,
            format="docx",
            width=794,
            height=1123,
            dpi=96,
            checksum=checksum,
            quality=spec.quality,
            task_status=(
                "degraded"
                if spec.quality.needs_review
                or spec.quality.degradation_reasons
                else "succeeded"
            ),
            file_name=file_name,
            render_manifest=render_manifest,
            non_empty_ratio=1.0,
            size_bytes=target.stat().st_size,
            generated_at=datetime.now(UTC),
            page_count=page_count,
            control_fields=dict(spec.control_fields),
        )

    def _replace_placeholders(
        self,
        document: DocumentObject,
        spec: DocumentRenderSpec,
    ) -> None:
        values = dict(spec.control_fields)
        values["file_name"] = "正式成果文件"
        values["page_number"] = "1"
        values["version"] = spec.template_version
        values["generated_by"] = "地震应急辅助决策系统"
        for paragraph in _iter_document_paragraphs(document):
            self._replace_paragraph_placeholders(paragraph, values)

    def _replace_paragraph_placeholders(
        self,
        paragraph: Any,
        values: Mapping[str, Any],
    ) -> None:
        if "{{" not in paragraph.text:
            return
        text = paragraph.text
        for key, value in values.items():
            text = text.replace(f"{{{{{key}}}}}", "" if value is None else str(value))
        paragraph.text = text

    def _remove_section_placeholders(self, document: DocumentObject) -> None:
        for paragraph in _iter_document_paragraphs(document):
            if (
                "{{section_1_title}}" in paragraph.text
                or "{{section_1_body}}" in paragraph.text
                or "{{section_1_image}}" in paragraph.text
            ):
                paragraph.text = ""

    def _append_sections(
        self,
        document: DocumentObject,
        spec: DocumentRenderSpec,
    ) -> None:
        for index, section in enumerate(spec.sections):
            if index > 0:
                document.add_page_break()
            document.add_heading(section.title, level=1)
            for line in section.body:
                document.add_paragraph(line)
            if section.table:
                self._add_table(document, section.table)

    def _add_table(self, document: DocumentObject, rows: Sequence[Sequence[str]]) -> None:
        table = document.add_table(rows=0, cols=max(len(row) for row in rows))
        table.style = "Table Grid"
        for row_values in rows:
            row = table.add_row()
            for index, value in enumerate(row_values):
                row.cells[index].text = str(value)

    def _embed_images(
        self,
        document: DocumentObject,
        spec: DocumentRenderSpec,
        output_parent: Path,
    ) -> None:
        images = self._all_images(spec, output_parent)
        if not images:
            return
        first = next(iter(images.values()))
        insertion = self._find_blank_placeholder_paragraph(document)
        if insertion is not None:
            self._add_image_to_paragraph(insertion, first)
        else:
            document.add_paragraph()
            self._add_image_to_paragraph(document.paragraphs[-1], first)
        for image in list(images.values())[1:]:
            paragraph = document.add_paragraph()
            self._add_image_to_paragraph(paragraph, image)

    def _add_image_to_paragraph(self, paragraph: Any, image: DocumentImage) -> None:
        paragraph.text = ""
        run = paragraph.add_run()
        run.add_picture(
            str(image.path),
            width=Cm(15),
        )
        if image.caption:
            paragraph.add_run(f" {image.caption}")

    def _find_blank_placeholder_paragraph(self, document: DocumentObject) -> Any | None:
        for paragraph in document.paragraphs:
            if paragraph.text.strip() == "":
                return paragraph
        return None

    def _set_footer(self, document: DocumentObject, spec: DocumentRenderSpec) -> None:
        footer = document.sections[0].footer
        footer.paragraphs[0].text = (
            "正式成果文件 | 第 1 页 | 版本 "
            f"{spec.template_version} | 地震应急辅助决策系统"
        )
        footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _assert_no_placeholders(self, document: DocumentObject) -> None:
        text = "\n".join(
            paragraph.text for paragraph in _iter_document_paragraphs(document)
        )
        if PLACEHOLDER_PATTERN.search(text):
            raise ValueError("rendered DOCX still contains unresolved placeholders")

    def _page_count(self, document: DocumentObject) -> int:
        count = 1
        for paragraph in document.paragraphs:
            if _paragraph_has_page_break(paragraph):
                count += 1
        return count

    def _all_images(
        self,
        spec: DocumentRenderSpec,
        output_parent: Path,
    ) -> dict[str, DocumentImage]:
        images = {image.key: image for image in spec.images}
        chart_root = output_parent / ".charts"
        for chart_spec in spec.charts:
            chart_path = chart_root / f"{spec.artifact_key}-{chart_spec.chart_id}.png"
            result = self._chart_renderer.render(chart_spec, chart_path)
            images[chart_spec.chart_id] = DocumentImage(
                key=chart_spec.chart_id,
                path=result.path,
                checksum=result.checksum,
                width=result.width,
                height=result.height,
                caption=chart_spec.title,
            )
        return images

    def _resolved_image_checksums(
        self,
        spec: DocumentRenderSpec,
        output_parent: Path,
    ) -> dict[str, str]:
        checksums: dict[str, str] = {}
        for key, image in self._all_images(spec, output_parent).items():
            checksum = image.checksum or _sha256_path(image.path)
            if image.checksum and checksum != image.checksum:
                raise ValueError(f"image checksum mismatch for {key}")
            checksums[key] = checksum
        return checksums


def _iter_document_paragraphs(document: DocumentObject) -> tuple[Any, ...]:
    paragraphs: list[Any] = list(document.paragraphs)
    for table in document.tables:
        paragraphs.extend(_iter_table_paragraphs(table))
    for section in document.sections:
        for container in (section.header, section.footer):
            paragraphs.extend(container.paragraphs)
            for table in container.tables:
                paragraphs.extend(_iter_table_paragraphs(table))
    return tuple(paragraphs)


def _iter_table_paragraphs(table: Any) -> tuple[Any, ...]:
    paragraphs: list[Any] = []
    for row in table.rows:
        for cell in row.cells:
            paragraphs.extend(cell.paragraphs)
            for nested_table in cell.tables:
                paragraphs.extend(_iter_table_paragraphs(nested_table))
    return tuple(paragraphs)


def _paragraph_has_page_break(paragraph: Any) -> bool:
    return paragraph._element.xpath(".//w:br[@w:type='page']")


def _image_dimensions(path: Path) -> tuple[int, int]:
    with Image.open(path) as image:
        return image.size
