from __future__ import annotations

import io
import os
from pathlib import Path
from typing import Any
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt as PptxPt

CONTROL_FIELDS = (
    ("event_name", "事件名称"),
    ("magnitude", "震级"),
    ("origin_time", "发震时刻"),
    ("longitude", "震中经度"),
    ("latitude", "震中纬度"),
    ("depth_km", "震源深度"),
    ("t1_at", "T1"),
    ("data_source", "数据来源"),
    ("report_type", "报告类型"),
    ("mode_marker", "成果标识"),
    ("generated_at", "生成时间"),
    ("artifact_version", "成果版本"),
    ("revision_no", "事件修订号"),
    ("assessment_run_no", "评估运行号"),
    ("data_asset_snapshot_fingerprint", "数据资产快照指纹"),
    ("model_parameter_version", "模型参数版本"),
    ("template_version", "模板版本"),
    ("quality_grade", "质量等级"),
    ("degradation_reasons", "降级原因"),
    ("needs_review", "待复核项"),
)


DECISION_TEMPLATE_SLIDES = (
    (
        "cover",
        "封面",
        (
            "辅助决策报告",
            "{{title}}",
            "{{event_name}}",
            "{{origin_time}}",
            "{{magnitude}}",
            "{{quality_grade}}",
            "{{mode_marker}}",
        ),
        (),
    ),
    (
        "event_facts",
        "事件参数",
        ("{{event_facts}}",),
        (),
    ),
    (
        "response_recommendation",
        "响应建议",
        ("{{recommendation}}",),
        (),
    ),
    (
        "fused_intensity",
        "融合烈度",
        ("{{intensity}}",),
        ("map.intensity",),
    ),
    (
        "casualties_building_damage",
        "伤亡与房屋破坏",
        ("{{casualties}}", "{{buildings}}"),
        ("map.deaths", "map.injuries", "map.building_damage"),
    ),
    (
        "economic_resource_demand",
        "经济与资源需求",
        ("{{economy}}", "{{resources}}"),
        ("map.economic_loss", "map.rescue_demand", "map.material_demand"),
    ),
    (
        "key_targets_faults",
        "重点目标与断裂",
        ("{{key_targets}}", "{{faults}}"),
        ("map.key_targets", "map.active_faults"),
    ),
    (
        "spatial_distance_conclusions",
        "空间距离与结论",
        (
            "{{spatial_conclusions}}",
            "{{source_versions}}",
            "{{degradation_reasons}}",
            "{{needs_review}}",
        ),
        ("map.epicenter", "map.city_distances", "map.buried"),
    ),
)


def build_background_templates(output_dir: Path) -> dict[str, Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / "background-template.docx"
    temporary = output_dir / f".background-template.{os.getpid()}.docx"

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.start_type = WD_SECTION.NEW_PAGE

    normal = document.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("{{title}}").bold = True
    title.runs[0].font.size = Pt(22)

    document.add_heading("事件控制表", level=1)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for field_name, label in CONTROL_FIELDS:
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = f"{{{{{field_name}}}}}"

    document.add_heading("{{section_1_title}}", level=1)
    document.add_paragraph("{{section_1_body}}")
    document.add_paragraph("{{section_1_image}}")

    footer = document.sections[0].footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.text = (
        "{{file_name}} | 第 {{page_number}} 页 | 版本 {{version}} | "
        "{{generated_by}}"
    )

    _save_deterministic_docx(document, temporary)
    os.replace(temporary, target)
    return {"background-template": target}


def build_decision_template(output_dir: Path) -> Path:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / "decision-template.pptx"
    temporary = output_dir / f".decision-template.{os.getpid()}.pptx"

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]

    for layout_key, layout_label, body_lines, image_keys in DECISION_TEMPLATE_SLIDES:
        slide = presentation.slides.add_slide(blank_layout)
        _add_decision_title(slide, layout_key, layout_label)

        top = Inches(1.45)
        if layout_key == "cover":
            top = Inches(2.0)
        for index, line in enumerate(body_lines):
            _add_decision_text_box(
                slide,
                line,
                left=Inches(0.9),
                top=top + Inches(index * 0.72),
                width=Inches(11.5),
                height=Inches(0.68),
                bold=layout_key == "cover" and index == 0,
                size=24 if layout_key == "cover" and index == 0 else 18,
                alignment=PP_ALIGN.CENTER if layout_key == "cover" else PP_ALIGN.LEFT,
            )

        if image_keys:
            _add_decision_image_slots(slide, image_keys)

    _save_deterministic_pptx(presentation, temporary)
    os.replace(temporary, target)
    return target


def _add_decision_title(slide: Any, layout_key: str, layout_label: str) -> None:
    box = slide.shapes.add_textbox(
        Inches(0.7),
        Inches(0.28),
        Inches(11.9),
        Inches(0.72),
    )
    box.name = f"decision-{layout_key}"
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = layout_label
    run.font.name = "Noto Sans CJK SC"
    run.font.size = PptxPt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(20, 30, 40)


def _add_decision_text_box(
    slide: Any,
    text: str,
    *,
    left: Any,
    top: Any,
    width: Any,
    height: Any,
    bold: bool = False,
    size: int = 18,
    alignment: Any = PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = alignment
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Noto Sans CJK SC"
    run.font.size = PptxPt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(20, 30, 40)


def _add_decision_image_slots(slide: Any, image_keys: tuple[str, ...]) -> None:
    columns = 3
    width = Inches(3.82)
    height = Inches(2.72)
    gap_x = Inches(0.18)
    gap_y = Inches(0.16)
    left = Inches(0.45)
    top = Inches(4.05)
    for index, image_key in enumerate(image_keys):
        row, column = divmod(index, columns)
        box = slide.shapes.add_textbox(
            left + column * (width + gap_x),
            top + row * (height + gap_y),
            width,
            height,
        )
        box.name = f"decision-image-{image_key}"
        frame = box.text_frame
        frame.word_wrap = True
        frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        run = frame.paragraphs[0].add_run()
        run.text = f"{{{{image:{image_key}}}}}"
        run.font.name = "Noto Sans CJK SC"
        run.font.size = PptxPt(16)
        run.font.color.rgb = RGBColor(120, 120, 120)


def _save_deterministic_docx(document: Document, target: Path) -> None:
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    with zipfile.ZipFile(buffer, "r") as source:
        with zipfile.ZipFile(
            target,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as destination:
            for item in source.infolist():
                info = zipfile.ZipInfo(
                    item.filename,
                    date_time=(1980, 1, 1, 0, 0, 0),
                )
                info.compress_type = zipfile.ZIP_DEFLATED
                destination.writestr(
                    info,
                    source.read(item.filename),
                )


def _save_deterministic_pptx(presentation: Presentation, target: Path) -> None:
    buffer = io.BytesIO()
    presentation.save(buffer)
    buffer.seek(0)
    with zipfile.ZipFile(buffer, "r") as source:
        with zipfile.ZipFile(
            target,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as destination:
            for item in source.infolist():
                info = zipfile.ZipInfo(
                    item.filename,
                    date_time=(1980, 1, 1, 0, 0, 0),
                )
                info.compress_type = zipfile.ZIP_DEFLATED
                destination.writestr(
                    info,
                    source.read(item.filename),
                )
