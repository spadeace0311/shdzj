from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
from PIL import Image, UnidentifiedImageError

from app.artifacts.domain import ArtifactDefinition, ArtifactKind, ProductionMode
from app.config import settings

_MAP_DIMENSIONS = (4761, 3369)
_MAP_DPI = 300
_PLACEHOLDER_PATTERN = re.compile(r"\{\{[^{}]*\}\}")


@dataclass(frozen=True, slots=True)
class ValidationResult:
    valid: bool
    error_category: str | None = None
    summary: str = ""
    checksum: str | None = None
    dimensions: tuple[int, int] | None = None
    page_count: int | None = None


class ArtifactValidator:
    def __init__(
        self,
        *,
        max_bytes: int | None = None,
        min_non_empty_pixel_ratio: float = 0.001,
    ) -> None:
        self._max_bytes = (
            settings.artifact_max_override_bytes
            if max_bytes is None
            else max_bytes
        )
        if self._max_bytes < 1:
            raise ValueError("max_bytes must be positive")
        if not 0 < min_non_empty_pixel_ratio < 1:
            raise ValueError("min_non_empty_pixel_ratio must be between 0 and 1")
        self._min_non_empty_pixel_ratio = min_non_empty_pixel_ratio

    def validate(
        self,
        path: Path,
        definition: ArtifactDefinition,
        production_mode: ProductionMode,
    ) -> ValidationResult:
        mode = (
            production_mode.value
            if isinstance(production_mode, ProductionMode)
            else str(production_mode)
        )
        source = Path(path)
        if not source.exists():
            return _invalid("file_missing", "artifact file does not exist")
        if not source.is_file():
            return _invalid("file_missing", "artifact path is not a regular file")
        size_bytes = source.stat().st_size
        if size_bytes <= 0:
            return _invalid("file_missing", "artifact file is empty")
        if size_bytes > self._max_bytes:
            return _invalid("size_exceeded", "artifact upload exceeds maximum size")

        checksum = _sha256_path(source)
        extension = _extension(source)
        expected_extension = definition.format.lower()
        if expected_extension == "jpg":
            allowed_extensions = {"jpg", "jpeg"}
        else:
            allowed_extensions = {expected_extension}
        if extension not in allowed_extensions:
            return _invalid(
                "format_mismatch",
                f"expected .{expected_extension} but found .{extension}",
                checksum=checksum,
            )
        if not _magic_matches(source, extension):
            return _invalid(
                "format_mismatch",
                "file magic bytes do not match the declared extension",
                checksum=checksum,
            )

        if definition.kind == ArtifactKind.MAP:
            return self._validate_map(
                source,
                checksum=checksum,
                expected_format=expected_extension,
                production_mode=mode,
            )
        if definition.kind == ArtifactKind.DOCX:
            return self._validate_docx(source, checksum=checksum)
        if definition.kind == ArtifactKind.PPTX:
            return self._validate_pptx(source, checksum=checksum)
        return _invalid("format_mismatch", "unsupported artifact kind")

    def _validate_map(
        self,
        path: Path,
        *,
        checksum: str,
        expected_format: str,
        production_mode: str,
    ) -> ValidationResult:
        expected_pillow_format = "JPEG" if expected_format == "jpg" else "PNG"
        try:
            with Image.open(path) as image:
                image.verify()
            with Image.open(path) as image:
                actual_format = image.format
                width, height = image.size
                dpi = image.info.get("dpi")
                marker_text = _image_marker_text(path, image)
                non_empty_ratio = _non_empty_pixel_ratio(image)
        except (UnidentifiedImageError, OSError, ValueError) as error:
            return _invalid(
                "format_mismatch",
                f"image parser could not read file: {error}",
                checksum=checksum,
            )

        if actual_format != expected_pillow_format:
            return _invalid(
                "format_mismatch",
                f"expected {expected_pillow_format} but parsed {actual_format}",
                checksum=checksum,
            )
        if (width, height) != _MAP_DIMENSIONS:
            return _invalid(
                "geometry_mismatch",
                f"expected {_MAP_DIMENSIONS} but found {(width, height)}",
                checksum=checksum,
                dimensions=(width, height),
            )
        if dpi is None or tuple(dpi) != (_MAP_DPI, _MAP_DPI):
            return _invalid(
                "dpi_mismatch",
                f"expected {_MAP_DPI} DPI but found {dpi}",
                checksum=checksum,
                dimensions=(width, height),
            )
        if non_empty_ratio < self._min_non_empty_pixel_ratio:
            return _invalid(
                "empty_map",
                "map contains too few non-empty pixels",
                checksum=checksum,
                dimensions=(width, height),
            )
        if production_mode in {"test", "drill", "replay"}:
            if not _contains_marker(marker_text):
                return _invalid(
                    "marker_missing",
                    "test, drill, or replay map must include a visible or metadata marker",
                    checksum=checksum,
                    dimensions=(width, height),
                )
        return ValidationResult(
            valid=True,
            summary="map format, geometry, DPI, and pixel coverage are valid",
            checksum=checksum,
            dimensions=(width, height),
        )

    def _validate_docx(
        self,
        path: Path,
        *,
        checksum: str,
    ) -> ValidationResult:
        try:
            from docx import Document

            document = Document(path)
            text = _docx_text(document)
        except Exception as error:
            return _invalid(
                "format_mismatch",
                f"DOCX parser could not read file: {error}",
                checksum=checksum,
            )
        if _PLACEHOLDER_PATTERN.search(text):
            return _invalid(
                "unresolved_placeholder",
                "DOCX contains unresolved placeholder markers",
                checksum=checksum,
            )
        return ValidationResult(
            valid=True,
            summary="DOCX opened and contains no unresolved placeholders",
            checksum=checksum,
            page_count=1,
        )

    def _validate_pptx(
        self,
        path: Path,
        *,
        checksum: str,
    ) -> ValidationResult:
        try:
            from pptx import Presentation

            presentation = Presentation(path)
            text = _pptx_text(presentation)
            page_count = len(presentation.slides)
        except Exception as error:
            return _invalid(
                "format_mismatch",
                f"PPTX parser could not read file: {error}",
                checksum=checksum,
            )
        if _PLACEHOLDER_PATTERN.search(text):
            return _invalid(
                "unresolved_placeholder",
                "PPTX contains unresolved placeholder markers",
                checksum=checksum,
                page_count=page_count,
            )
        return ValidationResult(
            valid=True,
            summary="PPTX opened and contains no unresolved placeholders",
            checksum=checksum,
            page_count=page_count,
        )


def _invalid(
    category: str,
    summary: str,
    *,
    checksum: str | None = None,
    dimensions: tuple[int, int] | None = None,
    page_count: int | None = None,
) -> ValidationResult:
    return ValidationResult(
        valid=False,
        error_category=category,
        summary=summary,
        checksum=checksum,
        dimensions=dimensions,
        page_count=page_count,
    )


def _extension(path: Path) -> str:
    return path.suffix.lower().lstrip(".")


def _magic_matches(path: Path, extension: str) -> bool:
    with path.open("rb") as source:
        head = source.read(16)
    if extension in {"jpg", "jpeg"}:
        return head.startswith(b"\xff\xd8\xff")
    if extension == "png":
        return head.startswith(b"\x89PNG\r\n\x1a\n")
    if extension in {"docx", "pptx"}:
        return head.startswith(b"PK\x03\x04")
    return False


def _non_empty_pixel_ratio(image: Image.Image) -> float:
    grayscale = np.asarray(image.convert("L"), dtype=np.uint8)
    non_empty = grayscale < 250
    return float(np.mean(non_empty)) if non_empty.size else 0.0


def _image_marker_text(path: Path, image: Image.Image) -> str:
    pieces: list[str] = [path.name]
    info = image.info or {}
    for key in ("comment", "Comment", "description", "Description"):
        value = info.get(key)
        if value:
            pieces.append(str(value))
    try:
        exif = image.getexif()
        if exif:
            for value in exif.values():
                pieces.append(str(value))
    except Exception:
        pass
    return "\n".join(pieces)


def _contains_marker(text: str) -> bool:
    return any(marker in text for marker in ("测试", "演练", "测试回放"))


def _docx_text(document: Any) -> str:
    pieces: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                pieces.append(cell.text)
    for section in document.sections:
        for container in (section.header, section.footer):
            if container is None:
                continue
            pieces.extend(paragraph.text for paragraph in container.paragraphs)
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        pieces.append(cell.text)
    return "\n".join(pieces)


def _pptx_text(presentation: Any) -> str:
    pieces: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                pieces.extend(
                    paragraph.text
                    for paragraph in shape.text_frame.paragraphs
                )
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        pieces.append(cell.text)
    return "\n".join(pieces)


def _sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while True:
            chunk = source.read(1024 * 1024)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()
