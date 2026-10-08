from __future__ import annotations

import csv
import io
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from openpyxl import load_workbook
from pypdf import PdfReader


class UnsupportedDocumentError(ValueError):
    """Raised when the document format is not supported."""


class UnreadableDocumentError(ValueError):
    """Raised when a document cannot be read or has no reliable text."""


@dataclass(frozen=True, slots=True)
class ParsedBlock:
    kind: str
    text: str
    page: int | None
    section_path: tuple[str, ...]
    row_range: tuple[int, int] | None
    metadata: dict[str, Any]


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    title: str
    blocks: list[ParsedBlock]
    metadata: dict[str, Any]


_HEADING_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
_HTML_BLOCK_TAGS = {"p", "li", "blockquote", "pre", "dt", "dd", "figcaption"}
_HTML_CONTAINER_TAGS = {
    "a",
    "article",
    "body",
    "div",
    "footer",
    "header",
    "html",
    "main",
    "ol",
    "section",
    "span",
    "ul",
}
_EXCLUDED_HTML_TAGS = {
    "embed",
    "iframe",
    "link",
    "meta",
    "noscript",
    "object",
    "script",
    "style",
    "template",
}


class DocumentParser:
    def parse(
        self,
        path: str | Path,
        *,
        file_name: str,
        mime_type: str | None = None,
    ) -> ParsedDocument:
        path_obj = Path(path)
        suffix = Path(file_name).suffix.lower()
        if not suffix:
            suffix = path_obj.suffix.lower()
        if suffix not in _HANDLERS and mime_type:
            suffix = _MIME_SUFFIXES.get(mime_type, suffix)

        handler = _HANDLERS.get(suffix)
        if handler is None:
            raise UnsupportedDocumentError(
                f"unsupported document format: {file_name}"
            )
        if not path_obj.exists():
            raise UnreadableDocumentError(
                f"document file does not exist: {file_name}"
            )

        try:
            title, blocks = handler(path_obj, file_name)
        except (UnreadableDocumentError, UnsupportedDocumentError):
            raise
        except Exception as exc:
            raise UnreadableDocumentError(
                f"unable to read {file_name}: {exc}"
            ) from exc

        metadata = {
            "file_name": file_name,
            "mime_type": mime_type,
            "format": suffix.lstrip("."),
        }
        return ParsedDocument(title=title, blocks=blocks, metadata=metadata)


def _parse_markdown_or_text(
    path: Path,
    file_name: str,
) -> tuple[str, list[ParsedBlock]]:
    text = _read_text(path, file_name)
    blocks: list[ParsedBlock] = []
    section_path: tuple[str, ...] = ()
    title: str | None = None
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_lines:
            return
        paragraph = "\n".join(paragraph_lines).strip()
        paragraph_lines.clear()
        if paragraph:
            blocks.append(
                ParsedBlock(
                    kind="paragraph",
                    text=paragraph,
                    page=None,
                    section_path=section_path,
                    row_range=None,
                    metadata={},
                )
            )

    for raw_line in text.splitlines():
        line = raw_line.strip()
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            heading_text = heading.group(2).strip()
            section_path = section_path[: level - 1] + (heading_text,)
            if title is None and level == 1:
                title = heading_text
            blocks.append(
                ParsedBlock(
                    kind="heading",
                    text=heading_text,
                    page=None,
                    section_path=section_path,
                    row_range=None,
                    metadata={},
                )
            )
            continue
        if line:
            paragraph_lines.append(line)
        else:
            flush_paragraph()
    flush_paragraph()
    return title or Path(file_name).stem, blocks


def _parse_html(
    path: Path,
    file_name: str,
) -> tuple[str, list[ParsedBlock]]:
    text = _read_text(path, file_name)
    soup = BeautifulSoup(text, "lxml")
    for tag in soup(_EXCLUDED_HTML_TAGS):
        tag.decompose()

    title = soup.title.get_text("", strip=True) if soup.title else None
    blocks: list[ParsedBlock] = []
    section_path: tuple[str, ...] = ()
    pending_text: list[str] = []
    _walk_html(soup.body or soup, section_path, blocks, pending_text)
    _flush_html_text(pending_text, section_path, blocks)
    return title or Path(file_name).stem, blocks


def _walk_html(
    element: Tag,
    section_path: tuple[str, ...],
    blocks: list[ParsedBlock],
    pending_text: list[str],
) -> None:
    for child in element.children:
        if isinstance(child, NavigableString):
            direct_text = str(child)
            if direct_text.strip():
                pending_text.append(direct_text)
            continue
        if not isinstance(child, Tag):
            continue
        name = child.name.lower()
        if name in _EXCLUDED_HTML_TAGS:
            continue
        if name in _HEADING_TAGS:
            _flush_html_text(pending_text, section_path, blocks)
            heading_text = child.get_text("", strip=True)
            if not heading_text:
                continue
            level = int(name[1])
            section_path = section_path[: level - 1] + (heading_text,)
            blocks.append(
                ParsedBlock(
                    kind="heading",
                    text=heading_text,
                    page=None,
                    section_path=section_path,
                    row_range=None,
                    metadata={},
                )
            )
            continue
        if name == "table":
            _flush_html_text(pending_text, section_path, blocks)
            table_block = _html_table_block(child, section_path)
            if table_block is not None:
                blocks.append(table_block)
            continue
        if name in _HTML_BLOCK_TAGS:
            _flush_html_text(pending_text, section_path, blocks)
            block_text = child.get_text("", strip=True)
            if block_text:
                blocks.append(
                    ParsedBlock(
                        kind="paragraph",
                        text=block_text,
                        page=None,
                        section_path=section_path,
                        row_range=None,
                        metadata={},
                    )
                )
            continue
        if name in _HTML_CONTAINER_TAGS:
            _walk_html(child, section_path, blocks, pending_text)


def _flush_html_text(
    pending_text: list[str],
    section_path: tuple[str, ...],
    blocks: list[ParsedBlock],
) -> None:
    if not pending_text:
        return
    text = re.sub(r"\s+", " ", "".join(pending_text)).strip()
    pending_text.clear()
    if text:
        blocks.append(
            ParsedBlock(
                kind="paragraph",
                text=text,
                page=None,
                section_path=section_path,
                row_range=None,
                metadata={},
            )
        )


def _html_table_block(
    table: Tag,
    section_path: tuple[str, ...],
) -> ParsedBlock | None:
    rows: list[list[str]] = []
    for tr in table.find_all("tr"):
        if tr.find_parent("table") is not table:
            continue
        cells = [
            cell.get_text("", strip=True)
            for cell in tr.find_all(["th", "td"])
            if cell.find_parent("table") is table
        ]
        if cells:
            rows.append(cells)
    if not rows:
        return None
    headers = rows[0]
    data_rows = rows[1:]
    row_range = (2, len(rows)) if data_rows else None
    metadata = {"headers": headers, "rows": data_rows}
    return ParsedBlock(
        kind="table",
        text=_serialize_table(headers, data_rows),
        page=None,
        section_path=section_path,
        row_range=row_range,
        metadata=metadata,
    )


def _parse_docx(
    path: Path,
    file_name: str,
) -> tuple[str, list[ParsedBlock]]:
    document = Document(str(path))
    blocks: list[ParsedBlock] = []
    section_path: tuple[str, ...] = ()
    title: str | None = None

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = DocxParagraph(child, document)
            paragraph_text = paragraph.text.strip()
            if not paragraph_text:
                continue
            level = _docx_heading_level(paragraph)
            if level is None:
                blocks.append(
                    ParsedBlock(
                        kind="paragraph",
                        text=paragraph_text,
                        page=None,
                        section_path=section_path,
                        row_range=None,
                        metadata={},
                    )
                )
                continue
            section_path = section_path[: level - 1] + (paragraph_text,)
            if title is None and level == 1:
                title = paragraph_text
            blocks.append(
                ParsedBlock(
                    kind="heading",
                    text=paragraph_text,
                    page=None,
                    section_path=section_path,
                    row_range=None,
                    metadata={},
                )
            )
            continue
        if isinstance(child, CT_Tbl):
            table = DocxTable(child, document)
            table_block = _docx_table_block(table, section_path)
            if table_block is not None:
                blocks.append(table_block)

    return title or Path(file_name).stem, blocks


def _docx_table_block(
    table: DocxTable,
    section_path: tuple[str, ...],
) -> ParsedBlock | None:
    rows = [
        [cell.text.strip() for cell in row.cells]
        for row in table.rows
    ]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return None
    headers = rows[0]
    data_rows = rows[1:]
    row_range = (2, len(rows)) if data_rows else None
    metadata = {"headers": headers, "rows": data_rows}
    return ParsedBlock(
        kind="table",
        text=_serialize_table(headers, data_rows),
        page=None,
        section_path=section_path,
        row_range=row_range,
        metadata=metadata,
    )


def _docx_heading_level(paragraph: DocxParagraph) -> int | None:
    style_name = (paragraph.style.name or "") if paragraph.style else ""
    digits = re.findall(r"\d+", style_name)
    if digits:
        return int(digits[0])
    if style_name.lower() in {"title", "标题"}:
        return 1
    return None


def _parse_pdf(
    path: Path,
    file_name: str,
) -> tuple[str, list[ParsedBlock]]:
    reader = PdfReader(str(path))
    blocks: list[ParsedBlock] = []
    text_characters = 0
    section_path: tuple[str, ...] = ()

    for page_number, page in enumerate(reader.pages, start=1):
        font_runs: list[tuple[str, float]] = []
        page_text = page.extract_text(
            visitor_text=lambda text, _cm, _tm, _font_dict, font_size,
            runs=font_runs: _record_pdf_font_run(runs, text, font_size)
        ) or ""
        text_characters += len(re.sub(r"\s+", "", page_text))
        page_blocks, section_path = _pdf_blocks_from_text(
            page_text,
            page_number=page_number,
            font_runs=font_runs,
            section_path=section_path,
        )
        blocks.extend(page_blocks)

    if text_characters < 20:
        raise UnreadableDocumentError(
            "PDF contains no extractable text; OCR is not supported"
        )
    return Path(file_name).stem, blocks


def _record_pdf_font_run(
    runs: list[tuple[str, float]],
    text: object,
    font_size: object,
) -> None:
    raw_text = "" if text is None else str(text)
    size = 0.0 if font_size is None else float(font_size)
    if raw_text.strip():
        runs.append((raw_text, size))


def _pdf_blocks_from_text(
    page_text: str,
    *,
    page_number: int,
    font_runs: list[tuple[str, float]],
    section_path: tuple[str, ...],
) -> tuple[list[ParsedBlock], tuple[str, ...]]:
    heading_model = _pdf_heading_model(font_runs)
    blocks: list[ParsedBlock] = []

    for raw_line in page_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        level = _pdf_heading_level(line, heading_model)
        if level is None:
            blocks.append(
                ParsedBlock(
                    kind="paragraph",
                    text=line,
                    page=page_number,
                    section_path=section_path,
                    row_range=None,
                    metadata={"page": page_number},
                )
            )
            continue
        section_path = section_path[: level - 1] + (line,)
        blocks.append(
            ParsedBlock(
                kind="heading",
                text=line,
                page=page_number,
                section_path=section_path,
                row_range=None,
                metadata={"page": page_number},
            )
        )
    return blocks, section_path


def _pdf_heading_model(
    runs: list[tuple[str, float]],
) -> dict[str, object]:
    fragments: dict[str, float] = {}
    size_counts: Counter[float] = Counter()
    for raw_text, size in runs:
        normalized = _normalize_pdf_text(raw_text)
        if not normalized or size <= 0:
            continue
        size_counts[size] += 1
        fragments[normalized] = max(fragments.get(normalized, 0.0), size)

    if not size_counts:
        return {}
    body_size = min(size_counts)
    threshold = max(body_size * 1.3, body_size + 2.0)
    heading_sizes = sorted(
        {size for size in size_counts if size >= threshold},
        reverse=True,
    )
    return {
        "fragments": fragments,
        "threshold": threshold,
        "level_by_size": {
            size: min(index + 1, 6)
            for index, size in enumerate(heading_sizes)
        },
    }


def _pdf_heading_level(
    line: str,
    model: dict[str, object],
) -> int | None:
    if not model or len(line) > 160:
        return None
    fragments = model["fragments"]
    level_by_size = model["level_by_size"]
    threshold = model["threshold"]
    if not isinstance(fragments, dict) or not isinstance(level_by_size, dict):
        return None
    normalized_line = _normalize_pdf_text(line)
    candidate_sizes: list[float] = []
    for fragment, size in fragments.items():
        if (
            size >= float(threshold)
            and size in level_by_size
            and (fragment == normalized_line or fragment in normalized_line)
        ):
            candidate_sizes.append(float(size))
    if not candidate_sizes:
        return None
    return int(level_by_size[max(candidate_sizes)])


def _normalize_pdf_text(text: str) -> str:
    return " ".join(text.split())


def _parse_xlsx(
    path: Path,
    file_name: str,
) -> tuple[str, list[ParsedBlock]]:
    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True,
    )
    blocks: list[ParsedBlock] = []
    for sheet in workbook.worksheets:
        rows: list[list[str]] = []
        for values in sheet.iter_rows(values_only=True):
            row = ["" if value is None else str(value) for value in values]
            if any(cell.strip() for cell in row):
                rows.append(row)
        if not rows:
            continue
        headers = rows[0]
        data_rows = rows[1:]
        row_range = (2, len(rows)) if data_rows else None
        metadata = {
            "headers": headers,
            "rows": data_rows,
            "sheet_name": sheet.title,
        }
        blocks.append(
            ParsedBlock(
                kind="table",
                text=_serialize_table(headers, data_rows),
                page=None,
                section_path=(sheet.title,),
                row_range=row_range,
                metadata=metadata,
            )
        )
    return Path(file_name).stem, blocks


def _parse_csv(
    path: Path,
    file_name: str,
) -> tuple[str, list[ParsedBlock]]:
    text = _read_text(path, file_name)
    rows: list[list[str]] = []
    for row in csv.reader(io.StringIO(text)):
        if any(cell.strip() for cell in row):
            rows.append([cell.strip() for cell in row])
    if not rows:
        return Path(file_name).stem, []
    headers = rows[0]
    data_rows = rows[1:]
    row_range = (2, len(rows)) if data_rows else None
    metadata = {"headers": headers, "rows": data_rows}
    block = ParsedBlock(
        kind="table",
        text=_serialize_table(headers, data_rows),
        page=None,
        section_path=(),
        row_range=row_range,
        metadata=metadata,
    )
    return Path(file_name).stem, [block]


def _read_text(path: Path, file_name: str) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnreadableDocumentError(
        f"document is not valid UTF-8 or GB18030 text: {file_name}"
    )


def _serialize_table(headers: list[str], rows: list[list[str]]) -> str:
    lines = []
    if headers:
        lines.append(" | ".join(headers))
    for row in rows:
        lines.append(" | ".join(row))
    return "\n".join(lines)


_MIME_SUFFIXES = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/xhtml+xml": ".html",
    "text/csv": ".csv",
    "text/html": ".html",
    "text/markdown": ".md",
    "text/plain": ".txt",
}

_HANDLERS: dict[str, Callable[[Path, str], tuple[str, list[ParsedBlock]]]] = {
    ".csv": _parse_csv,
    ".docx": _parse_docx,
    ".htm": _parse_html,
    ".html": _parse_html,
    ".md": _parse_markdown_or_text,
    ".pdf": _parse_pdf,
    ".txt": _parse_markdown_or_text,
    ".xlsx": _parse_xlsx,
}
