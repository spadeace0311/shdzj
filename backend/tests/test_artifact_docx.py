from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from app.artifacts.renderers.chart_renderer import ChartRenderer, ChartSeries, ChartSpec
from app.artifacts.renderers.docx_renderer import (
    DocxRenderer,
    build_background_document_spec,
)

BACKGROUND_DOCS = (
    "doc.background",
    "doc.housing",
    "doc.economy",
    "doc.population",
    "doc.key_targets",
    "doc.spatial_distances",
    "doc.area_overview",
    "doc.historical_catalog",
)


@pytest.fixture(autouse=True)
async def _dispose_engine_between_tests():
    from app.db import engine

    await engine.dispose()
    yield
    await engine.dispose()


def _sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while True:
            chunk = source.read(1024 * 1024)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


@pytest.fixture
def docx_renderer() -> DocxRenderer:
    return DocxRenderer()


@pytest.fixture
def chart_renderer() -> ChartRenderer:
    return ChartRenderer()


@pytest.mark.parametrize("artifact_key", BACKGROUND_DOCS)
async def test_background_documents_are_valid_docx(
    artifact_key: str,
    seeded_artifact_assessment,
    docx_renderer: DocxRenderer,
    tmp_path: Path,
) -> None:
    context = await seeded_artifact_assessment.document_context(artifact_key)
    result = await docx_renderer.render(
        build_background_document_spec(context, artifact_key),
        tmp_path / f"{artifact_key}.docx",
    )
    document = Document(result.path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert result.page_count >= 1
    assert "{{" not in text
    assert "事件名称" in text
    assert "数据来源" in text
    assert result.control_fields["template_version"] == "v1"


async def test_docx_contains_mode_marker_in_pixels_and_metadata(
    seeded_artifact_assessment,
    docx_renderer: DocxRenderer,
    tmp_path: Path,
) -> None:
    context = await seeded_artifact_assessment.document_context(
        "doc.background",
        production_mode="test",
    )
    result = await docx_renderer.render(
        build_background_document_spec(context, "doc.background"),
        tmp_path / "test.docx",
    )
    document = Document(result.path)

    assert "【测试】" in result.file_name
    assert "【测试】" in document.paragraphs[0].text
    assert result.render_manifest["marker"] == "【测试】"


async def test_docx_reuses_the_same_map_checksum_for_repeated_reference(
    seeded_artifact_assessment,
    docx_renderer: DocxRenderer,
    tmp_path: Path,
) -> None:
    context = await seeded_artifact_assessment.document_context("doc.area_overview")
    spec = build_background_document_spec(context, "doc.area_overview")
    result = await docx_renderer.render(spec, tmp_path / "overview.docx")

    assert result.render_manifest["image_checksums"]["map.epicenter"] == (
        context.artifacts["map.epicenter"].checksum
    )


async def test_docx_uses_not_applicable_t1_when_null(
    seeded_artifact_assessment,
    docx_renderer: DocxRenderer,
    tmp_path: Path,
) -> None:
    context = await seeded_artifact_assessment.document_context(
        "doc.background",
        t1_at=None,
    )
    result = await docx_renderer.render(
        build_background_document_spec(context, "doc.background"),
        tmp_path / "manual.docx",
    )
    text = "\n".join(paragraph.text for paragraph in Document(result.path).paragraphs)

    assert "T1: 不适用（人工/测试/演练）" in text
    assert "不适用（人工/测试/演练）" in text


async def test_chart_renderer_generates_deterministic_local_png(
    chart_renderer: ChartRenderer,
    tmp_path: Path,
) -> None:
    spec = ChartSpec(
        chart_id="population-age",
        title="常住人口年龄结构",
        labels=("0-14岁", "15-64岁", "65岁及以上"),
        series=(
            ChartSeries(name="常住人口", values=(18.2, 71.6, 10.2)),
        ),
        y_label="比例（%）",
    )
    first_path = tmp_path / "first.png"
    second_path = tmp_path / "second.png"

    first = chart_renderer.render(spec, first_path)
    second = chart_renderer.render(spec, second_path)

    assert first.checksum == second.checksum
    assert first.render_manifest["image_checksums"]["population-age"] == first.checksum
    with Image.open(first_path) as image:
        assert image.size == (spec.width, spec.height)
        assert image.format == "PNG"
        assert tuple(
            round(value) for value in image.info.get("dpi", (0, 0))
        ) == (spec.dpi, spec.dpi)


@pytest.mark.parametrize(
    ("artifact_key", "expected_sections"),
    (
        (
            "doc.background",
            {"事件参数", "所在行政区", "历史地震", "邻近断裂", "基础地理背景"},
        ),
        (
            "doc.housing",
            {"街镇房屋总量", "结构类型", "破坏统计", "覆盖与质量"},
        ),
        (
            "doc.economy",
            {"GDP", "产业结构", "经济损失", "来源与情景"},
        ),
        (
            "doc.population",
            {"常住人口", "流动人口", "家庭户", "年龄结构", "受灾人口"},
        ),
        (
            "doc.key_targets",
            {"避难场所", "学校", "医院", "危险源", "救援队伍", "文物单位", "重点目标列表"},
        ),
        (
            "doc.spatial_distances",
            {"城市距离", "区县距离", "街镇距离", "主要城市距离", "重点目标距离", "断裂距离"},
        ),
        (
            "doc.area_overview",
            {"地理概况", "行政区", "烈度", "人口", "房屋", "经济", "关键风险"},
        ),
        (
            "doc.historical_catalog",
            {"检索半径", "震级阈值", "历史地震目录", "灾害地震目录", "统计"},
        ),
    ),
)
async def test_each_background_document_has_its_required_sections(
    artifact_key: str,
    expected_sections: set[str],
    seeded_artifact_assessment,
    docx_renderer: DocxRenderer,
    tmp_path: Path,
) -> None:
    context = await seeded_artifact_assessment.document_context(artifact_key)
    spec = build_background_document_spec(context, artifact_key)
    result = await docx_renderer.render(spec, tmp_path / f"{artifact_key}.docx")
    text = "\n".join(paragraph.text for paragraph in Document(result.path).paragraphs)

    for section in expected_sections:
        assert section in text
