from pathlib import Path

import pytest

from app.knowledge.parser import (
    DocumentParser,
    UnreadableDocumentError,
    UnsupportedDocumentError,
)
from docx import Document as DocxDocument
from openpyxl import Workbook
from pypdf import PdfWriter


def test_markdown_parser_keeps_heading_path(tmp_path: Path) -> None:
    path = tmp_path / "doc.md"
    path.write_text("# 第一章\n\n## 总则\n\n震后立即报告。", encoding="utf-8")

    parsed = DocumentParser().parse(path, file_name="doc.md")

    assert parsed.blocks[-1].section_path == ("第一章", "总则")
    assert parsed.blocks[-1].text == "震后立即报告。"


def test_text_parser_falls_back_to_gb18030(tmp_path: Path) -> None:
    path = tmp_path / "doc.txt"
    path.write_bytes("应急响应启动。".encode("gb18030"))

    parsed = DocumentParser().parse(path, file_name="doc.txt")

    assert parsed.blocks[-1].text == "应急响应启动。"


def test_html_parser_removes_scripts_styles_and_links(tmp_path: Path) -> None:
    path = tmp_path / "doc.html"
    path.write_text(
        """
        <html>
          <head>
            <title>处置预案</title>
            <style>body { color: red; }</style>
            <script>alert("should not run")</script>
          </head>
          <body>
            <h1>第一章</h1>
            <p>震后立即报告。<a href="https://example.invalid/script">外部</a></p>
          </body>
        </html>
        """,
        encoding="utf-8",
    )

    parsed = DocumentParser().parse(path, file_name="doc.html")
    body_text = "\n".join(block.text for block in parsed.blocks if block.text)

    assert parsed.title == "处置预案"
    assert "alert" not in body_text
    assert "body { color" not in body_text
    assert parsed.blocks[-1].text == "震后立即报告。外部"
    assert parsed.blocks[-1].section_path == ("第一章",)


def test_csv_parser_uses_first_row_as_header(tmp_path: Path) -> None:
    path = tmp_path / "doc.csv"
    path.write_text(
        "区域,响应等级\n浦东,一级\n徐汇,二级\n",
        encoding="utf-8",
    )

    parsed = DocumentParser().parse(path, file_name="doc.csv")
    table = next(block for block in parsed.blocks if block.kind == "table")

    assert table.metadata["headers"] == ["区域", "响应等级"]
    assert table.metadata["rows"] == [["浦东", "一级"], ["徐汇", "二级"]]
    assert table.row_range == (2, 3)


def test_xlsx_parser_converts_each_sheet_to_header_rows(
    tmp_path: Path,
) -> None:
    path = tmp_path / "doc.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "响应"
    sheet.append(["区域", "等级"])
    sheet.append(["浦东", "一级"])
    sheet.append(["徐汇", "二级"])
    workbook.save(path)

    parsed = DocumentParser().parse(path, file_name="doc.xlsx")
    table = next(block for block in parsed.blocks if block.kind == "table")

    assert table.metadata["sheet_name"] == "响应"
    assert table.metadata["headers"] == ["区域", "等级"]
    assert table.metadata["rows"] == [["浦东", "一级"], ["徐汇", "二级"]]


def test_docx_parser_keeps_heading_and_table(tmp_path: Path) -> None:
    path = tmp_path / "doc.docx"
    document = DocxDocument()
    document.add_heading("第一章", level=1)
    document.add_paragraph("震后立即报告。")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "区域"
    table.rows[0].cells[1].text = "等级"
    table.rows[1].cells[0].text = "浦东"
    table.rows[1].cells[1].text = "一级"
    document.save(path)

    parsed = DocumentParser().parse(path, file_name="doc.docx")
    paragraph = next(block for block in parsed.blocks if block.kind == "paragraph")
    table_block = next(block for block in parsed.blocks if block.kind == "table")

    assert paragraph.section_path == ("第一章",)
    assert table_block.metadata["headers"] == ["区域", "等级"]
    assert table_block.metadata["rows"] == [["浦东", "一级"]]


def test_image_only_pdf_fails_without_ocr(tmp_path: Path) -> None:
    path = tmp_path / "image-only.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    with path.open("wb") as target:
        writer.write(target)

    with pytest.raises(UnreadableDocumentError, match="extractable text"):
        DocumentParser().parse(path, file_name="image-only.pdf")


def test_unknown_extension_is_unsupported(tmp_path: Path) -> None:
    path = tmp_path / "doc.bin"
    path.write_bytes(b"not a supported document")

    with pytest.raises(UnsupportedDocumentError):
        DocumentParser().parse(path, file_name="doc.bin")
