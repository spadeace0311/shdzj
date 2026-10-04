from __future__ import annotations

from io import BytesIO

import pytest
from PIL import Image, ImageDraw

from app.artifacts.domain import ArtifactDefinition, ArtifactKind, ProductionMode
from app.artifacts.storage import ArtifactStore
from app.artifacts.validation import ArtifactValidator


def _map_definition(*, file_format: str = "jpg") -> ArtifactDefinition:
    return ArtifactDefinition(
        artifact_key="map.epicenter",
        display_name="震中位置分布图",
        kind=ArtifactKind.MAP,
        output_profile="a3v-professional",
        format=file_format,
        legacy_code="M26",
        priority=80,
        depends_on=(),
        optional_depends_on=(),
        required_assets=(),
        optional_assets=(),
        template_key="map.epicenter",
        marker_policy="mode",
        failure_policy="block",
        quality_policy="A",
    )


def _docx_definition() -> ArtifactDefinition:
    return ArtifactDefinition(
        artifact_key="doc.rapid_brief",
        display_name="快速评估简报",
        kind=ArtifactKind.DOCX,
        output_profile="a3v-professional",
        format="docx",
        legacy_code="D09",
        priority=90,
        depends_on=(),
        optional_depends_on=(),
        required_assets=(),
        optional_assets=(),
        template_key="doc.rapid_brief",
        marker_policy="mode",
        failure_policy="block",
        quality_policy="A",
    )


def _pptx_definition() -> ArtifactDefinition:
    return ArtifactDefinition(
        artifact_key="deck.decision_report",
        display_name="辅助决策报告",
        kind=ArtifactKind.PPTX,
        output_profile="a3v-professional",
        format="pptx",
        legacy_code="D12",
        priority=90,
        depends_on=(),
        optional_depends_on=(),
        required_assets=(),
        optional_assets=(),
        template_key="deck.decision_report",
        marker_policy="mode",
        failure_policy="block",
        quality_policy="A",
    )


def _write_map_jpeg(path, *, comment: str | None = None) -> None:
    image = Image.new("RGB", (4761, 3369), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((80, 80, 1200, 1200), fill=(255, 0, 0))
    image.save(
        path,
        format="JPEG",
        quality=10,
        dpi=(300, 300),
        comment=comment,
    )


def test_store_is_content_addressed_and_rejects_path_escape(tmp_path) -> None:
    store = ArtifactStore(tmp_path, max_override_bytes=1024)
    stored = store.store_immutable_stream(
        BytesIO(b"artifact-bytes"),
        file_name="map.jpg",
    )

    assert (
        stored.checksum
        == "6521df166eb07efaf36eba5b6bedefd9d6a252e9c80bab1c99653700ec71473c"
    )
    assert stored.relative_path.endswith("-map.jpg")
    with pytest.raises(ValueError):
        store.resolve("../outside.jpg")


def test_store_immutable_stream_cleans_staging(tmp_path) -> None:
    store = ArtifactStore(tmp_path, max_override_bytes=1024)
    store.store_immutable_stream(BytesIO(b"artifact-bytes"), file_name="map.jpg")

    staging = tmp_path / "staging"
    assert staging.exists()
    assert list(staging.iterdir()) == []


def test_delete_unreferenced_removes_object(tmp_path) -> None:
    store = ArtifactStore(tmp_path, max_override_bytes=1024)
    stored = store.store_immutable_stream(
        BytesIO(b"artifact-bytes"),
        file_name="map.jpg",
    )

    assert stored.managed_path.exists()
    store.delete_unreferenced(stored)
    assert not stored.managed_path.exists()


def test_store_rejects_oversized_stream_and_cleans_staging(tmp_path) -> None:
    store = ArtifactStore(tmp_path, max_override_bytes=4)

    with pytest.raises(ValueError):
        store.store_immutable_stream(
            BytesIO(b"too-large"),
            file_name="map.jpg",
        )

    assert list((tmp_path / "staging").iterdir()) == []


def test_resolve_rejects_absolute_and_backslash_escapes(tmp_path) -> None:
    store = ArtifactStore(tmp_path, max_override_bytes=1024)

    for candidate in ("/etc/passwd", "..\\outside.jpg", "C:\\outside.jpg"):
        with pytest.raises(ValueError):
            store.resolve(candidate)


def test_validator_rejects_extension_mismatch(tmp_path) -> None:
    path = tmp_path / "fake.jpg"
    path.write_bytes(b"not-an-image")

    result = ArtifactValidator().validate(
        path,
        _map_definition(),
        production_mode=ProductionMode.LIVE,
    )

    assert result.valid is False
    assert result.error_category == "format_mismatch"


def test_validator_rejects_map_with_wrong_dimensions(tmp_path) -> None:
    path = tmp_path / "small.jpg"
    path.write_bytes(b"\xff\xd8\xff\xe0not-a-real-jpeg")

    result = ArtifactValidator().validate(
        path,
        _map_definition(),
        production_mode=ProductionMode.LIVE,
    )

    assert result.valid is False
    assert result.error_category in {"format_mismatch", "geometry_mismatch"}


def test_validator_docx_accepts_clean_document_and_rejects_placeholder(
    tmp_path,
) -> None:
    from docx import Document

    clean = tmp_path / "clean.docx"
    document = Document()
    document.add_paragraph("无占位符正文")
    document.save(clean)
    clean_result = ArtifactValidator().validate(
        clean,
        _docx_definition(),
        production_mode=ProductionMode.LIVE,
    )
    assert clean_result.valid is True
    assert clean_result.page_count == 1

    unresolved = tmp_path / "unresolved.docx"
    document = Document()
    document.add_paragraph("{{name}}")
    document.save(unresolved)
    unresolved_result = ArtifactValidator().validate(
        unresolved,
        _docx_definition(),
        production_mode=ProductionMode.LIVE,
    )
    assert unresolved_result.valid is False
    assert unresolved_result.error_category == "unresolved_placeholder"


def test_validator_docx_rejects_parser_failure(tmp_path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"PK\x03\x04not-a-docx")

    result = ArtifactValidator().validate(
        path,
        _docx_definition(),
        production_mode=ProductionMode.LIVE,
    )

    assert result.valid is False
    assert result.error_category == "format_mismatch"


def test_validator_docx_requires_exact_mode_marker(tmp_path) -> None:
    from docx import Document

    wrong_marker = tmp_path / "wrong.docx"
    document = Document()
    document.add_paragraph("【演练】")
    document.save(wrong_marker)
    wrong_result = ArtifactValidator().validate(
        wrong_marker,
        _docx_definition(),
        production_mode=ProductionMode.TEST,
    )
    assert wrong_result.valid is False
    assert wrong_result.error_category == "marker_missing"

    correct_marker = tmp_path / "correct.docx"
    document = Document()
    document.add_paragraph("【测试】")
    document.save(correct_marker)
    correct_result = ArtifactValidator().validate(
        correct_marker,
        _docx_definition(),
        production_mode=ProductionMode.TEST,
    )
    assert correct_result.valid is True


def test_validator_docx_requires_replay_marker(tmp_path) -> None:
    from docx import Document

    wrong_marker = tmp_path / "wrong-replay.docx"
    document = Document()
    document.add_paragraph("【测试】")
    document.save(wrong_marker)
    wrong_result = ArtifactValidator().validate(
        wrong_marker,
        _docx_definition(),
        production_mode=ProductionMode.REPLAY,
    )
    assert wrong_result.valid is False
    assert wrong_result.error_category == "marker_missing"

    correct_marker = tmp_path / "correct-replay.docx"
    document = Document()
    document.add_paragraph("【测试回放】")
    document.save(correct_marker)
    correct_result = ArtifactValidator().validate(
        correct_marker,
        _docx_definition(),
        production_mode=ProductionMode.REPLAY,
    )
    assert correct_result.valid is True


def test_validator_pptx_accepts_clean_deck_and_reports_page_count(
    tmp_path,
) -> None:
    from pptx import Presentation

    path = tmp_path / "clean.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(path)

    result = ArtifactValidator().validate(
        path,
        _pptx_definition(),
        production_mode=ProductionMode.LIVE,
    )

    assert result.valid is True
    assert result.page_count == 2


def test_validator_pptx_rejects_placeholder_and_parser_failure(tmp_path) -> None:
    from pptx import Presentation

    unresolved = tmp_path / "unresolved.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "{{name}}"
    presentation.save(unresolved)
    unresolved_result = ArtifactValidator().validate(
        unresolved,
        _pptx_definition(),
        production_mode=ProductionMode.LIVE,
    )
    assert unresolved_result.valid is False
    assert unresolved_result.error_category == "unresolved_placeholder"

    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"PK\x03\x04not-a-pptx")
    broken_result = ArtifactValidator().validate(
        broken,
        _pptx_definition(),
        production_mode=ProductionMode.LIVE,
    )
    assert broken_result.valid is False
    assert broken_result.error_category == "format_mismatch"


def test_validator_pptx_requires_exact_mode_marker(tmp_path) -> None:
    from pptx import Presentation

    wrong_marker = tmp_path / "wrong.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "【测试】"
    presentation.save(wrong_marker)
    wrong_result = ArtifactValidator().validate(
        wrong_marker,
        _pptx_definition(),
        production_mode=ProductionMode.DRILL,
    )
    assert wrong_result.valid is False
    assert wrong_result.error_category == "marker_missing"

    correct_marker = tmp_path / "correct.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "【演练】"
    presentation.save(correct_marker)
    correct_result = ArtifactValidator().validate(
        correct_marker,
        _pptx_definition(),
        production_mode=ProductionMode.DRILL,
    )
    assert correct_result.valid is True


def test_validator_pptx_requires_replay_marker(tmp_path) -> None:
    from pptx import Presentation

    wrong_marker = tmp_path / "wrong-replay.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "【测试】"
    presentation.save(wrong_marker)
    wrong_result = ArtifactValidator().validate(
        wrong_marker,
        _pptx_definition(),
        production_mode=ProductionMode.REPLAY,
    )
    assert wrong_result.valid is False
    assert wrong_result.error_category == "marker_missing"

    correct_marker = tmp_path / "correct-replay.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "【测试回放】"
    presentation.save(correct_marker)
    correct_result = ArtifactValidator().validate(
        correct_marker,
        _pptx_definition(),
        production_mode=ProductionMode.REPLAY,
    )
    assert correct_result.valid is True


def test_validator_map_ignores_file_name_marker_and_accepts_metadata(
    tmp_path,
) -> None:
    file_name_only = tmp_path / "【测试】fake.jpg"
    _write_map_jpeg(file_name_only, comment=None)
    file_name_result = ArtifactValidator().validate(
        file_name_only,
        _map_definition(),
        production_mode=ProductionMode.TEST,
    )
    assert file_name_result.valid is False
    assert file_name_result.error_category == "marker_missing"

    metadata_marker = tmp_path / "metadata.jpg"
    _write_map_jpeg(metadata_marker, comment="【测试】")
    metadata_result = ArtifactValidator().validate(
        metadata_marker,
        _map_definition(),
        production_mode=ProductionMode.TEST,
    )
    assert metadata_result.valid is True

    wrong_metadata_marker = tmp_path / "wrong-metadata.jpg"
    _write_map_jpeg(wrong_metadata_marker, comment="【演练】")
    wrong_result = ArtifactValidator().validate(
        wrong_metadata_marker,
        _map_definition(),
        production_mode=ProductionMode.TEST,
    )
    assert wrong_result.valid is False
    assert wrong_result.error_category == "marker_missing"


def test_validator_map_requires_replay_metadata_marker(tmp_path) -> None:
    wrong_marker = tmp_path / "wrong-replay-map.jpg"
    _write_map_jpeg(wrong_marker, comment="【测试】")
    wrong_result = ArtifactValidator().validate(
        wrong_marker,
        _map_definition(),
        production_mode=ProductionMode.REPLAY,
    )
    assert wrong_result.valid is False
    assert wrong_result.error_category == "marker_missing"

    correct_marker = tmp_path / "correct-replay-map.jpg"
    _write_map_jpeg(correct_marker, comment="【测试回放】")
    correct_result = ArtifactValidator().validate(
        correct_marker,
        _map_definition(),
        production_mode=ProductionMode.REPLAY,
    )
    assert correct_result.valid is True
