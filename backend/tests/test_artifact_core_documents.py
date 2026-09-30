from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from app.artifacts.basemap import MapViewportTileManifest, VIEWPORT_RADII_KM
from app.artifacts.context import ProductionContextService
from app.artifacts.repository import ArtifactProductionRepository
from app.artifacts.renderers.docx_renderer import (
    DocxRenderer,
    build_core_document_spec,
)
from app.artifacts.renderers.pptx_renderer import PptxRenderer
from app.config import settings
from tests.basemap_fixtures import make_in_memory_package


@pytest.fixture(autouse=True)
async def _dispose_engine_between_tests():
    from app.db import engine

    await engine.dispose()
    yield
    await engine.dispose()


@pytest.fixture
def docx_renderer() -> DocxRenderer:
    return DocxRenderer()


@pytest.fixture
def pptx_renderer() -> PptxRenderer:
    return PptxRenderer()


@pytest.fixture
def production_context_service() -> ProductionContextService:
    manifests = tuple(
        MapViewportTileManifest.build(
            center_lon=121.5,
            center_lat=31.2,
            radius_km=radius_km,
            output_width=settings.artifact_basemap_output_width,
            output_height=settings.artifact_basemap_output_height,
            zoom_levels=settings.artifact_basemap_zoom_levels,
            padding=settings.artifact_basemap_buffer_pixels,
        )
        for radius_km in VIEWPORT_RADII_KM
    )
    return ProductionContextService(
        repository=ArtifactProductionRepository(),
        offline_basemap_packages={
            "gaode": make_in_memory_package(manifests, provider="gaode"),
            "tianditu": make_in_memory_package(manifests, provider="tianditu"),
        },
    )


def _docx_text(path: Path) -> str:
    return "\n".join(paragraph.text for paragraph in Document(path).paragraphs)


def _pptx_text(path: Path) -> str:
    presentation = Presentation(path)
    return "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )


def _sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while True:
            chunk = source.read(1024 * 1024)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


async def test_rapid_brief_contains_t1_not_applicable_for_manual_event(
    seeded_artifact_assessment,
    docx_renderer: DocxRenderer,
    tmp_path: Path,
) -> None:
    context = await seeded_artifact_assessment.document_context(
        "doc.rapid_brief",
        t1_at=None,
    )
    result = await docx_renderer.render(
        build_core_document_spec(context, "doc.rapid_brief"),
        tmp_path / "brief.docx",
    )
    text = _docx_text(result.path)

    assert "T1" in text
    assert "不适用（人工/测试/演练）" in text


async def test_rapid_report_omits_failed_optional_map_but_marks_degraded(
    seeded_artifact_assessment,
    docx_renderer: DocxRenderer,
    tmp_path: Path,
) -> None:
    context = await seeded_artifact_assessment.document_context(
        "doc.rapid_report",
        failed_optional_artifacts={"map.reservoirs"},
    )
    result = await docx_renderer.render(
        build_core_document_spec(context, "doc.rapid_report"),
        tmp_path / "report.docx",
    )
    text = _docx_text(result.path)

    assert result.task_status == "degraded"
    assert "map.reservoirs" not in result.render_manifest["images"]
    assert "水库数据待复核" in text


async def test_pptx_replaces_all_placeholders_and_uses_fixed_pages(
    seeded_artifact_assessment,
    pptx_renderer: PptxRenderer,
    tmp_path: Path,
) -> None:
    context = await seeded_artifact_assessment.document_context("deck.decision_report")
    result = await pptx_renderer.render(
        build_core_document_spec(context, "deck.decision_report"),
        tmp_path / "decision.pptx",
    )
    presentation = Presentation(result.path)
    all_text = _pptx_text(result.path)

    assert len(presentation.slides) == context.pptx_slide_count
    assert "{{" not in all_text
    assert "辅助决策报告" in all_text
    assert result.render_manifest["marker"] in {"", "【测试】", "【演练】", "【测试回放】"}


async def test_core_document_fails_when_hard_map_dependency_is_missing(
    seeded_artifact_assessment,
    docx_renderer: DocxRenderer,
    tmp_path: Path,
) -> None:
    context = await seeded_artifact_assessment.document_context(
        "doc.rapid_brief",
        missing_artifact_keys={"map.epicenter"},
    )

    with pytest.raises(FileNotFoundError, match="map.epicenter"):
        await docx_renderer.render(
            build_core_document_spec(context, "doc.rapid_brief"),
            tmp_path / "brief.docx",
        )


async def test_pptx_fails_when_hard_map_dependency_is_missing(
    seeded_artifact_assessment,
    pptx_renderer: PptxRenderer,
    tmp_path: Path,
) -> None:
    context = await seeded_artifact_assessment.document_context(
        "deck.decision_report",
        missing_artifact_keys={"map.epicenter"},
    )

    with pytest.raises(FileNotFoundError, match="map.epicenter"):
        await pptx_renderer.render(
            build_core_document_spec(context, "deck.decision_report"),
            tmp_path / "decision.pptx",
        )


async def test_pptx_manifest_tracks_images_and_checksums(
    seeded_artifact_assessment,
    pptx_renderer: PptxRenderer,
    tmp_path: Path,
) -> None:
    context = await seeded_artifact_assessment.document_context("deck.decision_report")
    spec = build_core_document_spec(context, "deck.decision_report")
    result = await pptx_renderer.render(spec, tmp_path / "decision.pptx")

    assert result.render_manifest["page_count"] == 8
    assert result.render_manifest["template_version"] == spec.template_version
    assert result.render_manifest["unresolved_placeholder_count"] == 0
    for key, checksum in result.render_manifest["image_checksums"].items():
        path = Path(result.render_manifest["image_paths"][key])
        assert _sha256_path(path) == checksum


@pytest.mark.parametrize(
    ("production_mode", "marker"),
    (
        ("test", "【测试】"),
        ("drill", "【演练】"),
        ("replay", "【测试回放】"),
    ),
)
async def test_core_pptx_mode_marker_is_in_file_name_body_and_manifest(
    production_mode: str,
    marker: str,
    seeded_artifact_assessment,
    pptx_renderer: PptxRenderer,
    tmp_path: Path,
) -> None:
    context = await seeded_artifact_assessment.document_context(
        "deck.decision_report",
        production_mode=production_mode,
    )
    result = await pptx_renderer.render(
        build_core_document_spec(context, "deck.decision_report"),
        tmp_path / f"{production_mode}.pptx",
    )
    presentation = Presentation(result.path)

    assert marker in result.file_name
    assert marker in result.render_manifest["marker"]
    assert any(marker in shape.text for shape in presentation.slides[0].shapes if hasattr(shape, "text"))


def test_decision_template_is_deterministic_and_has_no_linked_media(tmp_path: Path) -> None:
    from app.artifacts.template_builder import build_decision_template

    first = build_decision_template(tmp_path / "first")
    second = build_decision_template(tmp_path / "second")

    assert _sha256_path(first) == _sha256_path(second)
    presentation = Presentation(first)
    assert len(presentation.slides) == 8
    assert round(presentation.slide_width.cm) == 34
    assert round(presentation.slide_height.cm) == 19

    with zipfile.ZipFile(first) as archive:
        names = set(archive.namelist())
    assert not any(name.startswith("ppt/media/") for name in names)
    assert not any("externalLink" in name for name in names)


async def test_real_production_context_resolves_core_document_inputs(
    seeded_artifact_assessment,
    production_context_service,
    session_factory,
    tmp_path: Path,
) -> None:
    run = await seeded_artifact_assessment.create_full_run()
    async with session_factory() as session:
        async with session.begin():
            await production_context_service.freeze_static_context(
                session,
                run.id,
                seeded_artifact_assessment.catalog,
            )

    async with session_factory() as session:
        async with session.begin():
            task = await seeded_artifact_assessment.first_task(
                run.id,
                "doc.rapid_report",
            )
            context = await production_context_service.build_document_context(
                session,
                task.id,
                "docx",
            )

    spec = build_core_document_spec(context, "doc.rapid_report")
    assert spec.artifact_key == "doc.rapid_report"
    assert spec.template_version == "v1"
    assert spec.template_path.name == "background-template.docx"
    assert spec.quality.grade == "core"
